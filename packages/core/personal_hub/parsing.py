"""Разбор документов в текст (per-page).

Поддерживает TXT/MD (plain), PDF (pypdf + OCR пустых страниц), DOCX
(python-docx) и изображения (PNG/JPG/TIFF/WEBP — OCR). Возвращает
`ParsedDocument` (страницы + warnings). Ошибки OCR поднимаются как OcrError,
если текст не удалось извлечь вовсе; частичный сбой страниц виден в warnings
(документ индексируется по успешным страницам).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from personal_hub.ocr import OcrPagesResult

SUPPORTED_EXTENSIONS = {
    "txt", "md", "markdown", "pdf", "docx",
    "png", "jpg", "jpeg", "tif", "tiff", "webp", "bmp",
}

_IMAGE_MIMES = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/webp": "webp",
    "image/tiff": "tiff",
    "image/bmp": "bmp",
}


@dataclass
class TextPage:
    """Текст одной страницы/секции документа."""

    page: int
    text: str


@dataclass
class ParsedDocument:
    """Результат разбора: страницы с текстом и предупреждения.

    `warnings` заполняется при частичном OCR (часть страниц не распознана) —
    документ индексируется, но код предупреждения (например, OCR_TIMEOUT
    с номером страницы) виден в статусе/задаче.
    """

    pages: list[TextPage]
    warnings: list[str] = field(default_factory=list)


class UnsupportedFormatError(Exception):
    """Формат документа не поддерживается."""

    def __init__(self, fmt: str) -> None:
        super().__init__(f"неподдерживаемый формат документа: {fmt}")
        self.code = "UNSUPPORTED_FORMAT"


def _detect_format(filename: str, mime_type: str) -> str:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext in SUPPORTED_EXTENSIONS:
        return ext
    if mime_type == "application/pdf":
        return "pdf"
    if mime_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    }:
        return "docx"
    if mime_type in _IMAGE_MIMES:
        return _IMAGE_MIMES[mime_type]
    return ext


def parse_document(filename: str, mime_type: str, data: bytes) -> ParsedDocument:
    """Разбирает документ в страницы с текстом и предупреждениями."""
    fmt = _detect_format(filename, mime_type)
    if fmt == "pdf":
        return _parse_pdf(data)
    if fmt == "docx":
        return ParsedDocument(pages=_parse_docx(data))
    if fmt in {"txt", "md", "markdown"}:
        return ParsedDocument(pages=[TextPage(page=1, text=data.decode("utf-8", errors="replace"))])
    if fmt in SUPPORTED_EXTENSIONS and fmt not in {"txt", "md", "markdown", "pdf", "docx"}:
        return ParsedDocument(pages=_parse_image(data))
    raise UnsupportedFormatError(fmt)


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    texts: list[str] = []
    empty_pages: list[int] = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        texts.append(text)
        if not text:
            empty_pages.append(index + 1)  # 1-based

    # Для сканированных (без текстового слоя) страниц — OCR. Частичный сбой
    # страниц не отменяет разбор: он виден в warnings.
    warnings: list[str] = []
    if empty_pages:
        ocr = _ocr_pages(data, empty_pages)
        for page_num, text in ocr.texts.items():
            texts[page_num - 1] = text
        warnings = [
            f"{code} (стр. {page})" for page, code in sorted(ocr.failed.items())
        ]

    pages: list[TextPage] = []
    for index, text in enumerate(texts, start=1):
        if text.strip():
            pages.append(TextPage(page=index, text=text))
    return ParsedDocument(
        pages=pages or [TextPage(page=1, text="")], warnings=warnings
    )


def _ocr_pages(data: bytes, page_indices: list[int]) -> OcrPagesResult:
    """OCR заданных страниц PDF; результат содержит тексты и коды сбоев."""
    from personal_hub.ocr import ocr_pdf_pages

    return ocr_pdf_pages(data, page_indices)


def _parse_image(data: bytes) -> list[TextPage]:
    """OCR одиночного изображения документа (отдельный page).

    Защита: при отключённом OCR (`OCR_MAX_PAGES=0`) изображение не принимается;
    пиксельный лимит проверяется до `load()` (decompression bomb).
    """
    from PIL import Image

    from personal_hub.config import get_settings
    from personal_hub.ocr import OcrError, parse_languages, recognize_image

    settings = get_settings()
    if settings.ocr_max_pages == 0:
        raise OcrError("OCR_DISABLED", "OCR отключён (OCR_MAX_PAGES=0).")
    image = Image.open(io.BytesIO(data))
    width, height = image.size
    if width * height > settings.ocr_max_image_pixels:
        raise OcrError(
            "OCR_IMAGE_TOO_LARGE",
            f"Изображение слишком большое: {width}x{height} px "
            f"(лимит {settings.ocr_max_image_pixels} px).",
        )
    image.load()
    result = recognize_image(image, parse_languages(settings.ocr_languages))
    return [TextPage(page=1, text=result.text)]


def _parse_docx(data: bytes) -> list[TextPage]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Таблицы тоже извлекаем (простой вариант — по строкам).
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts)
    return [TextPage(page=1, text=text)]
